import shutil
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from investor_agent.models import MovementDraft  # noqa: E402
from investor_agent.parser import LocalMovementParser  # noqa: E402
from investor_agent.service import InvestmentAgentService  # noqa: E402


class DisabledAI:
    available = False


class RecordingMovementAI:
    available = True

    def __init__(self):
        self.messages = []

    def parse(self, text, base=None):
        self.messages.append(text)
        draft = MovementDraft.from_dict(base.to_dict() if base else None)
        draft.movement_type = "Compra"
        draft.trade_date = "2026-07-29"
        draft.ticker = "GOOGL"
        draft.quantity = 0.902
        draft.unit_price = 332.98
        return draft


class FakeResearcher:
    available = True

    def __init__(self):
        self.last_context = None
        self.last_include_advice = None

    def research(self, query, research_type, portfolio_context, include_advice=False):
        self.last_context = portfolio_context
        self.last_include_advice = include_advice
        return {
            "response_id": "resp_test",
            "model": "gpt-test",
            "answer": "## Resumen\nCandidata para investigar, no una compra segura.",
            "citations": [
                {
                    "url": "https://example.com/report",
                    "title": "Informe",
                    "start_index": 3,
                    "end_index": 10,
                }
            ],
            "sources": [{"url": "https://example.com/report", "title": "Informe"}],
            "search_queries": [query],
            "input_tokens": 10,
            "output_tokens": 20,
            "status": "completed",
            "incomplete_reason": None,
            "include_advice": include_advice,
            "duration_ms": 1234,
        }


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        temp_path = Path(self.temp.name)
        template = temp_path / "template.xlsx"
        shutil.copy2(ROOT / "assets" / "plantilla_base.xlsx", template)
        self.workbook = temp_path / "movimientos.xlsx"
        self.researcher = FakeResearcher()
        self.service = InvestmentAgentService(
            database_path=temp_path / "inversiones.db",
            template_path=template,
            workbook_path=self.workbook,
            local_parser=LocalMovementParser(today_provider=lambda: date(2026, 8, 13)),
            ai_parser=DisabledAI(),
            researcher=self.researcher,
        )

    def tearDown(self):
        self.temp.cleanup()

    def test_requires_confirmation_then_writes_db_and_excel(self):
        draft = self.service.handle_message(
            "s1", "Compré 3 AAPL a 180 USD en Wallbit con 1 de comisión"
        )
        self.assertEqual("confirmation", draft["kind"])
        self.assertEqual(0, self.service.db.stats()["movements"])

        saved = self.service.handle_message("s1", "confirmar")
        self.assertEqual("saved", saved["kind"])
        self.assertTrue(saved["excel_synced"])
        self.assertEqual(1, self.service.db.stats()["movements"])

        workbook = load_workbook(self.workbook, data_only=False)
        sheet = workbook["Movimientos"]
        self.assertEqual("Compra", sheet["B6"].value)
        self.assertEqual("AAPL", sheet["D6"].value)
        self.assertEqual(3, sheet["E6"].value)
        self.assertEqual(180, sheet["F6"].value)

    def test_missing_fields_can_be_completed_in_second_message(self):
        first = self.service.handle_message("s2", "Vendí 2 Microsoft")
        self.assertEqual("missing", first["kind"])
        self.assertCountEqual(
            ["account", "unit_price", "currency"], first["missing_fields"]
        )

        second = self.service.handle_message("s2", "Fue a 425 USD en Wallbit")
        self.assertEqual("confirmation", second["kind"])
        self.assertEqual("MSFT", second["draft"]["ticker"])
        self.assertEqual(425, second["draft"]["unit_price"])

    def test_cancel_does_not_save(self):
        self.service.handle_message("s3", "Compré 1 TSLA a 200 USD en Wallbit")
        cancelled = self.service.handle_message("s3", "cancelar")
        self.assertEqual("cancelled", cancelled["kind"])
        self.assertEqual(0, self.service.db.stats()["movements"])
        self.assertIsNone(self.service.db.get_pending("s3"))

    def test_defaults_and_positions(self):
        self.service.update_settings(
            {"default_account": "Wallbit", "default_currency": "USD"}
        )
        self.service.handle_message("s4", "Compré 5 KO a 60")
        self.service.handle_message("s4", "confirmar")
        positions = self.service.db.positions()
        self.assertEqual([{"ticker": "KO", "quantity": 5.0}], positions)

    def test_portfolio_word_does_not_bypass_the_llm_movement_parser(self):
        ai = RecordingMovementAI()
        self.service.ai_parser = ai
        self.service.update_settings(
            {"default_account": "Wallbit", "default_currency": "USD"}
        )

        result = self.service.handle_message(
            "llm-routing",
            "Agregar a mi cartera 0,902 acciones de GOOGL compradas el 29 de julio a 332,98 por accion",
        )

        self.assertEqual("confirmation", result["kind"])
        self.assertEqual("openai", result["parser_mode"])
        self.assertEqual(1, len(ai.messages))
        self.assertEqual("GOOGL", result["draft"]["ticker"])
        self.assertEqual(0.902, result["draft"]["quantity"])
        self.assertEqual(332.98, result["draft"]["unit_price"])

    def test_pending_draft_can_be_corrected_before_confirmation(self):
        self.service.update_settings(
            {"default_account": "Wallbit", "default_currency": "USD"}
        )
        first = self.service.handle_message(
            "correct-draft",
            "Agregar a mi cartera 0,902 acciones de GOOGL compradas el 29 de julio a 332,98 por accion",
        )
        self.assertEqual("confirmation", first["kind"])

        corrected = self.service.handle_message(
            "correct-draft",
            "El ticker correcto es GOOG y el precio correcto es 330,50",
        )

        self.assertEqual("confirmation", corrected["kind"])
        self.assertEqual("GOOG", corrected["draft"]["ticker"])
        self.assertEqual(330.50, corrected["draft"]["unit_price"])
        self.assertEqual(0.902, corrected["draft"]["quantity"])
        self.assertIn("Actualicé el borrador", corrected["message"])
        self.assertIn("Podés corregir cualquier dato", corrected["message"])
        self.assertEqual(0, self.service.db.stats()["movements"])

    def test_confirmed_movement_is_replaced_by_an_audited_amendment(self):
        self.service.update_settings(
            {"default_account": "Wallbit", "default_currency": "USD"}
        )
        self.service.db.upsert_cash_balance(
            as_of="2026-08-12",
            account="Wallbit",
            currency="USD",
            amount=1271,
            source_text="usuario",
        )
        first = self.service.handle_message(
            "amend-confirmed",
            "Compre 1200 USD en NVDA, 5,32 acciones a 224,81 por accion",
        )
        self.assertEqual(5.32, first["draft"]["quantity"])
        saved = self.service.handle_message("amend-confirmed", "confirmar")
        original_id = saved["movement_id"]

        correction = self.service.handle_message(
            "amend-confirmed", "Corrige mis acciones en NVDA, son 5,327"
        )
        self.assertEqual("confirmation", correction["kind"])
        self.assertEqual([original_id], correction["draft"]["amends_movement_ids"])
        self.assertEqual(5.327, correction["draft"]["quantity"])
        self.assertIn("enmienda auditada", correction["message"])

        amended = self.service.handle_message("amend-confirmed", "confirmar")
        self.assertEqual("saved", amended["kind"])
        self.assertNotEqual(original_id, amended["movement_id"])
        self.assertEqual([{"ticker": "NVDA", "quantity": 5.327}], self.service.db.positions())
        self.assertEqual(1, self.service.db.stats()["movements"])
        self.assertEqual(1, self.service.db.stats()["voided_movements"])
        self.assertIsNotNone(self.service.db.get_movement(original_id)["voided_at"])
        self.assertEqual(-1200, self.service.db.get_movement(original_id)["cash_delta_applied"])
        self.assertEqual(-1200, self.service.db.get_movement(amended["movement_id"])["cash_delta_applied"])
        self.assertEqual(71, self.service.db.latest_cash_balances()[0]["amount"])

        workbook = load_workbook(self.workbook, data_only=False)
        sheet = workbook["Movimientos"]
        self.assertEqual("Anulado", sheet["B6"].value)
        self.assertEqual("Compra", sheet["B7"].value)
        self.assertEqual(5.327, sheet["E7"].value)

    def test_confirmed_purchase_and_sale_update_available_cash(self):
        self.service.update_settings(
            {"default_account": "Wallbit", "default_currency": "USD"}
        )
        self.service.db.upsert_cash_balance(
            as_of="2026-08-12",
            account="Wallbit",
            currency="USD",
            amount=1000,
            source_text="usuario",
        )

        self.service.handle_message(
            "cash-purchase", "Compré 2 AAPL a 100 USD en Wallbit con 1 de comisión"
        )
        purchase = self.service.handle_message("cash-purchase", "confirmar")
        self.assertEqual("saved", purchase["kind"])
        self.assertIn("799,00 USD", purchase["message"])
        self.assertEqual(799, self.service.db.latest_cash_balances()[0]["amount"])

        self.service.handle_message(
            "cash-sale", "Vendí 1 AAPL a 120 USD en Wallbit con 2 de comisión"
        )
        sale = self.service.handle_message("cash-sale", "confirmar")
        self.assertEqual("saved", sale["kind"])
        self.assertIn("917,00 USD", sale["message"])
        self.assertEqual(917, self.service.db.latest_cash_balances()[0]["amount"])

    def test_import_is_idempotent_and_cash_is_separate(self):
        from investor_agent.models import MovementDraft

        draft = MovementDraft(
            movement_type="Compra",
            trade_date="2026-08-12",
            account="Wallbit",
            ticker="ASTI",
            quantity=43.95762389,
            unit_price=150 / 43.95762389,
            cash_amount=150,
            currency="USDC",
            source_text="captura",
        )
        first_id, first_created = self.service.db.import_movement(draft, "external-1")
        second_id, second_created = self.service.db.import_movement(draft, "external-1")
        self.assertTrue(first_created)
        self.assertFalse(second_created)
        self.assertEqual(first_id, second_id)
        self.assertEqual(1, self.service.db.stats()["movements"])

        self.service.db.upsert_cash_balance(
            as_of="2026-08-13",
            account="Wallbit",
            currency="USD",
            amount=1271,
            source_text="usuario",
        )
        balances = self.service.db.latest_cash_balances()
        self.assertEqual(1, len(balances))
        self.assertEqual(1271, balances[0]["amount"])
        self.assertEqual("USD", balances[0]["currency"])

    def test_research_is_saved_with_sources_and_portfolio_context(self):
        self.service.update_settings(
            {"default_account": "Wallbit", "default_currency": "USD"}
        )
        result = self.service.research(
            "research-session", "Investigá Visa", "asset"
        )

        self.assertEqual("research", result["kind"])
        self.assertEqual(1, result["id"])
        self.assertEqual("https://example.com/report", result["sources"][0]["url"])
        self.assertEqual(1, self.service.db.stats()["research_reports"])
        self.assertEqual("Wallbit", self.researcher.last_context["account"])
        self.assertIsNone(self.researcher.last_context["risk_limits"])
        self.assertEqual("completed", result["status"])
        self.assertEqual(1234, result["duration_ms"])

    def test_research_can_request_an_orientative_opinion(self):
        result = self.service.research(
            "advice-session",
            "¿Conviene comprar Visa?",
            "asset",
            include_advice=True,
        )

        self.assertTrue(result["include_advice"])
        self.assertTrue(result["context"]["advice_requested"])
        self.assertTrue(self.researcher.last_include_advice)

    def test_chat_detects_an_explicit_advice_request(self):
        result = self.service.handle_message(
            "chat-advice", "Investigá Visa y decime si conviene comprar"
        )

        self.assertEqual("research", result["kind"])
        self.assertTrue(result["include_advice"])

    def test_research_request_in_chat_uses_researcher(self):
        result = self.service.handle_message("chat-research", "Investigá Visa con fuentes")
        self.assertEqual("research", result["kind"])
        self.assertEqual("asset", result["research_type"])


if __name__ == "__main__":
    unittest.main()
