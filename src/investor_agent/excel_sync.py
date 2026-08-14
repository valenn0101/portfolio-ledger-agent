from __future__ import annotations

import shutil
import re
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any


class ExcelSync:
    def __init__(self, template_path: str | Path, workbook_path: str | Path):
        self.template_path = Path(template_path)
        self.workbook_path = Path(workbook_path)

    def ensure_workbook(self) -> Path:
        if self.workbook_path.exists():
            return self.workbook_path
        self.workbook_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.template_path, self.workbook_path)

        from openpyxl import load_workbook

        workbook = load_workbook(self.workbook_path)
        sheet = workbook["Movimientos"]
        for row in range(6, sheet.max_row + 1):
            for column in range(1, 11):
                sheet.cell(row, column).value = None
        workbook.save(self.workbook_path)
        return self.workbook_path

    def append(self, movement: dict[str, Any]) -> int:
        from openpyxl import load_workbook
        from openpyxl.formula.translate import Translator

        self.ensure_workbook()
        workbook = load_workbook(self.workbook_path)
        sheet = workbook["Movimientos"]
        row = self._first_empty_row(sheet)
        if row > sheet.max_row:
            source_row = max(6, row - 1)
            sheet.insert_rows(row)
            self._copy_row_layout(sheet, source_row, row, Translator)

        trade_date = datetime.strptime(movement["trade_date"], "%Y-%m-%d")
        values = (
            trade_date,
            movement["movement_type"],
            movement["account"],
            movement.get("ticker"),
            movement.get("quantity"),
            movement.get("unit_price"),
            movement.get("cash_amount"),
            movement["currency"],
            movement.get("fee"),
            self._notes(movement),
        )
        for column, value in enumerate(values, 1):
            sheet.cell(row, column).value = value
        sheet.cell(row, 1).number_format = "dd/mm/yyyy"

        self._expand_tables(sheet, row)
        workbook.save(self.workbook_path)
        return row

    def append_amendment(
        self, movement: dict[str, Any], voided_movement_ids: list[int]
    ) -> int:
        """Marca antecedentes como anulados y agrega su reemplazo auditado."""

        from openpyxl import load_workbook

        self.ensure_workbook()
        workbook = load_workbook(self.workbook_path)
        sheet = workbook["Movimientos"]
        missing_ids = set(int(value) for value in voided_movement_ids)
        for row in range(6, sheet.max_row + 1):
            notes = str(sheet.cell(row, 10).value or "")
            for movement_id in list(missing_ids):
                if re.search(rf"\bID DB:\s*{movement_id}\b", notes):
                    original_type = sheet.cell(row, 2).value
                    sheet.cell(row, 2).value = "Anulado"
                    sheet.cell(row, 10).value = (
                        f"{notes} | ANULADO por enmienda ID DB: {movement['id']} "
                        f"(tipo original: {original_type})"
                    )
                    missing_ids.remove(movement_id)
        if missing_ids:
            missing = ", ".join(f"#{value}" for value in sorted(missing_ids))
            raise RuntimeError(f"No encontré en Excel los movimientos a anular: {missing}")
        workbook.save(self.workbook_path)
        return self.append(movement)

    @staticmethod
    def _first_empty_row(sheet) -> int:
        for row in range(6, sheet.max_row + 1):
            if all(sheet.cell(row, column).value in (None, "") for column in range(1, 11)):
                return row
        return sheet.max_row + 1

    @staticmethod
    def _copy_row_layout(sheet, source_row: int, target_row: int, translator) -> None:
        sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
        for column in range(1, sheet.max_column + 1):
            source = sheet.cell(source_row, column)
            target = sheet.cell(target_row, column)
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            if source.alignment:
                target.alignment = copy(source.alignment)
            if source.protection:
                target.protection = copy(source.protection)
            if source.value is not None and isinstance(source.value, str) and source.value.startswith("="):
                target.value = translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)

    @staticmethod
    def _expand_tables(sheet, row: int) -> None:
        for table in sheet.tables.values():
            start, end = table.ref.split(":")
            end_column = "".join(character for character in end if character.isalpha())
            end_row = int("".join(character for character in end if character.isdigit()))
            if row > end_row:
                table.ref = f"{start}:{end_column}{row}"

    @staticmethod
    def _notes(movement: dict[str, Any]) -> str:
        parts = [movement.get("notes") or "Registrado por el agente"]
        if movement.get("inferred_fields"):
            parts.append("Inferido: " + ", ".join(movement["inferred_fields"]))
        parts.append(f"ID DB: {movement['id']}")
        return " | ".join(parts)
